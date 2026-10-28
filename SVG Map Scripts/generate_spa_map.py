import argparse
import json
import os
import re
import sys
from collections import OrderedDict

import openpyxl

# ── CONFIG ──────────────────────────────────────────────────────
EXCEL_NAME   = "RHTP SPA Commitments.xlsx"
OUTPUT_NAME  = "svg-map-spa-rhtp.json"
SEARCH_DEPTH = 2

COLOR_YES      = "046791"   # teal: has a commitment
COLOR_NO       = "F9CF41"   # yellow: no commitments
COLOR_HOVER    = "17415F"
NO_COMMIT_TEXT = "No State Policy Action Commitments"
BULLET         = "\u2022 "        # •
SUB_BULLET     = "    \u25e6 "    # ◦ with 4-space indent

# Words that identify an optional category/header column
CATEGORY_WORDS = ("category", "header", "description",
                  "factor name", "title", "group")

TAG_RE = re.compile(r"<[^>]+>")
TSF_STRIP_RE = re.compile(
    r"\bTechnical Score Factor\b[\s:\u2013-]*", re.I)
SPA_RE   = re.compile(r"\bSPA\b")    # whole word, uppercase
SPA_FULL = "State Policy Action"
CODE_RE  = re.compile(r"([A-Z])\.\s*(\d+)", re.I)


# ── HELPERS ─────────────────────────────────────────────────────
def clean(v):
    """Collapse whitespace and strip; None -> ''."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()

def norm_code(v):
    """'B.2.' / 'Technical Score Factor b. 2' -> 'B.2'"""
    m = CODE_RE.search(clean(v))
    if not m:
        return ""
    return f"{m.group(1).upper()}.{m.group(2)}"

def code_sort_key(code):
    letter, _, num = code.partition(".")
    return (letter, int(num) if num.isdigit() else 0)

def is_yes(v):
    return clean(v).lower() == "yes"

def tidy(text):
    """Drop 'Technical Score Factor', spell out SPA,
    keep any leading indent."""
    if not isinstance(text, str):
        return text
    stripped = text.lstrip(" ")
    indent = text[: len(text) - len(stripped)]
    body = TSF_STRIP_RE.sub("", stripped)
    body = SPA_RE.sub(SPA_FULL, body)
    body = re.sub(r" {2,}", " ", body).strip()
    return indent + body


# ── FIND COLUMNS ────────────────────────────────────────────────
def find_columns(cells):
    """cells = lowercased header row. Returns dict or None."""
    def first(test, skip=()):
        return next((i for i, c in enumerate(cells)
                     if c and i not in skip and test(c)),
                    None)

    category = first(lambda c: any(w in c
                                   for w in CATEGORY_WORDS))
    skip = () if category is None else (category,)

    col = {
        "state":  first(lambda c: c == "state", skip),
        "factor": first(lambda c: "factor" in c, skip),
        "item":   first(lambda c: "item" in c
                        or "policy" in c, skip),
        "commit": first(lambda c: "commit" in c
                        or "status" in c, skip),
    }
    if None in col.values():
        return None
    col["category"] = category      # may be None
    return col

def locate_table(wb):
    """Return (worksheet, header_row_index, columns)."""
    for ws in wb.worksheets:
        rows = ws.iter_rows(min_row=1, max_row=20,
                            values_only=True)
        for i, row in enumerate(rows):
            cells = [clean(c).lower() for c in row]
            col = find_columns(cells)
            if col:
                return ws, i, col
    return None, None, None


# ── READ DATA ───────────────────────────────────────────────────
def read_data(ws, header_idx, col):
    """
    Returns:
      data        {state: {code: [yes items]}}
      items       {code: [distinct items, in order]}
      categories  {code: category text from sheet}
    """
    data = OrderedDict()
    items = OrderedDict()
    categories = {}
    rows = ws.iter_rows(min_row=header_idx + 2,
                        values_only=True)
    for row in rows:
        cells = [clean(c) for c in row]

        def get(key):
            i = col.get(key)
            if i is None or i >= len(cells):
                return ""
            return cells[i]

        state = get("state")
        code = norm_code(get("factor"))
        item = tidy(get("item"))
        if not state or not code:
            continue

        cat = tidy(get("category"))
        if cat and code not in categories:
            categories[code] = cat

        code_items = items.setdefault(code, [])
        if item and item not in code_items:
            code_items.append(item)

        state_dict = data.setdefault(state, OrderedDict())
        if is_yes(get("commit")):
            state_dict.setdefault(code, []).append(item)
    return data, items, categories


# ── BUILD HOVER TEXT ────────────────────────────────────────────
def build_text(commitments, order, items, categories):
    lines = []
    for code in order:
        yes_items = commitments.get(code)
        if not yes_items:
            continue
        cat = categories.get(code, "")

        if len(items.get(code, [])) > 1:      # multi-item
            label = f"{BULLET}{code}: {cat}" if cat \
                    else f"{BULLET}{code}"
            lines.append(label + ":")
            for it in yes_items:
                lines.append(f"{SUB_BULLET}{it}")
        else:                                 # single item
            text = cat or yes_items[0]
            lines.append(f"{BULLET}{code}: {text}")
    return lines


# ── FILE DISCOVERY ──────────────────────────────────────────────
def search_dirs():
    dirs = []
    here = os.path.dirname(os.path.abspath(__file__))
    for d in (here, os.getcwd()):
        d = os.path.abspath(d)
        if d not in dirs:
            dirs.append(d)
    return dirs

def walk_files(extensions):
    seen = set()
    for root_dir in search_dirs():
        base = root_dir.rstrip(os.sep).count(os.sep)
        for root, subdirs, files in os.walk(root_dir):
            if root.count(os.sep) - base >= SEARCH_DEPTH:
                subdirs[:] = []
            subdirs[:] = [d for d in subdirs
                          if not d.startswith(".")]
            for name in files:
                path = os.path.abspath(
                    os.path.join(root, name))
                if path in seen or name.startswith("~$"):
                    continue
                if name.lower().endswith(extensions):
                    seen.add(path)
                    yield path

def choose(candidates, label):
    if not candidates:
        return None
    candidates.sort(key=os.path.getmtime, reverse=True)
    if len(candidates) == 1:
        return candidates[0]
    print(f"\nFound {len(candidates)} possible "
          f"{label} files:")
    for i, p in enumerate(candidates, 1):
        print(f"  {i}. {p}")
    if not sys.stdin.isatty():
        print(f"Using the most recent: {candidates[0]}")
        return candidates[0]
    while True:
        ans = input(f"Which {label}? "
                    f"[1-{len(candidates)}, Enter = 1]: ")
        ans = ans.strip()
        if not ans:
            return candidates[0]
        if ans.isdigit() and 1 <= int(ans) <= len(candidates):
            return candidates[int(ans) - 1]
        print("  Please enter a number from the list.")

def has_table(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws, _, _ = locate_table(wb)
        wb.close()
        return ws is not None
    except Exception:
        return False

def find_excel():
    files = list(walk_files((".xlsx", ".xlsm")))
    named = [p for p in files
             if os.path.basename(p).lower()
             == EXCEL_NAME.lower()]
    if named:
        return choose(named, f"'{EXCEL_NAME}'")
    return choose([p for p in files if has_table(p)],
                  "Excel")

def json_base_score(path):
    if os.path.basename(path).lower() == OUTPUT_NAME.lower():
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            doc = json.load(f)
    except Exception:
        return None
    if not isinstance(doc, dict):
        return None
    states = doc.get("stateData")
    if not isinstance(states, list) or not states:
        return None
    tagged = any("<u>" in str(s.get("stateName", ""))
                 for s in states)
    return 2 if tagged else 1

def find_base_json():
    scored = [(p, json_base_score(p))
              for p in walk_files((".json",))]
    scored = [(p, s) for p, s in scored if s is not None]
    if not scored:
        return None
    best = max(s for _, s in scored)
    return choose([p for p, s in scored if s == best],
                  "base JSON")


# ── MAIN ────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(
        description="Update State Policy Action map JSON.")
    ap.add_argument("--excel", help="Excel file (optional)")
    ap.add_argument("--json", help="Base JSON (optional)")
    ap.add_argument("--output", help="Output JSON (optional)")
    args = ap.parse_args()

    excel_file = args.excel or find_excel()
    if not excel_file or not os.path.exists(excel_file):
        print(f"ERROR: Could not find '{EXCEL_NAME}' in "
              f"{', '.join(search_dirs())} or subfolders.")
        sys.exit(1)

    json_file = args.json or find_base_json()
    if not json_file or not os.path.exists(json_file):
        print("ERROR: No base map JSON (with 'stateData') "
              f"found in {', '.join(search_dirs())}.")
        sys.exit(1)

    output_file = args.output or os.path.join(
        os.path.dirname(json_file), OUTPUT_NAME)

    print(f"\nExcel:     {excel_file}")
    print(f"Base JSON: {json_file}")
    print(f"Output:    {output_file}\n")

    try:
        wb = openpyxl.load_workbook(excel_file,
                                    data_only=True)
    except Exception as e:
        print(f"ERROR: Could not open Excel file: {e}")
        sys.exit(1)

    ws, header_idx, col = locate_table(wb)
    if ws is None:
        print("ERROR: Couldn't find a header row with "
              "State / Score Factor / Policy Item / "
              "Commitment columns. First rows of each tab:")
        for sheet in wb.worksheets:
            print(f"\n=== {sheet.title} ===")
            for row in sheet.iter_rows(min_row=1, max_row=6,
                                       values_only=True):
                print(row)
        sys.exit(1)

    data, items, categories = read_data(ws, header_idx, col)
    order = sorted(items, key=code_sort_key)

    print(f"Reading tab '{ws.title}' "
          f"(header on row {header_idx + 1})")
    if col["category"] is None:
        print("No Category column found (optional).")
    print("Factors found:")
    missing = []
    for code in order:
        multi = len(items[code]) > 1
        cat = categories.get(code)
        if cat:
            source = f"{cat}  [Category column]"
        elif not multi:
            source = f"{items[code][0]}  [Policy Item]"
        else:
            source = "(code only - no category in sheet)"
            missing.append(code)
        kind = "multi" if multi else "single"
        print(f"  {code}: {source} ({kind})")
        if multi:
            for it in items[code]:
                print(f"      - {it}")
    if missing:
        print(f"\nNote: {', '.join(missing)} have several "
              "items but no category name in the sheet, so "
              "they'll show as just the code. Add a "
              "'Category' column to give them a header.")

    with open(json_file, "r", encoding="utf-8") as fh:
        map_json = json.load(fh)

    updated, untouched, no_commit = 0, [], []
    json_states = set()
    for entry in map_json.get("stateData", []):
        plain = TAG_RE.sub("", entry.get("stateName", ""))
        plain = plain.strip()
        json_states.add(plain)
        if plain not in data:      # DC, territories
            untouched.append(plain)
            continue

        lines = build_text(data[plain], order, items,
                           categories)
        color = COLOR_YES if lines else COLOR_NO
        if not lines:
            no_commit.append(plain)
        entry["initialStateColor"] = color
        entry["stateOverColor"] = COLOR_HOVER
        entry["stateSelectedColor"] = color
        entry["text"] = lines or [NO_COMMIT_TEXT]
        updated += 1

    # Final cleanup on every entry (incl. DC/territories)
    for entry in map_json.get("stateData", []):
        if isinstance(entry.get("text"), list):
            entry["text"] = [tidy(t) for t in entry["text"]]
    side = map_json.get("defaultSideText")
    if isinstance(side, str):
        map_json["defaultSideText"] = tidy(side)

    with open(output_file, "w", encoding="utf-8") as fh:
        json.dump(map_json, fh, indent=2,
                  ensure_ascii=False)

    not_in_json = sorted(set(data) - json_states)
    print(f"\nDone! Updated {updated} states.")
    print(f"No commitments ({len(no_commit)}): "
          f"{', '.join(no_commit) or 'none'}")
    if untouched:
        print("Left unchanged (not in spreadsheet): "
              f"{', '.join(untouched)}")
    if not_in_json:
        print("WARNING: in spreadsheet but not in JSON: "
              f"{', '.join(not_in_json)}")
    print(f"Output: {output_file}")


if __name__ == "__main__":
    main()
