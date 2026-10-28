"""
generate_spa_table.py

Reads the one-tab RHTP SPA Commitments spreadsheet
(State | Score Factor | Policy Item | Active Commitment)
and builds a DataTable JSON file: {"aaData": [ {...}, ... ]}

Output columns:
  1. State
  2. State Policy Action Commitments  (count of "Yes",
     or "None" when there are none)
  3+ One column per Policy Item, "Yes" or "No"

Usage:
  python3 generate_spa_table.py
  python3 generate_spa_table.py -i "RHTP SPA Commitments.xlsx"
                                -o spa-table.json
"""

import argparse
import difflib
import json
import os
import re
import sys
from collections import OrderedDict

import openpyxl

# ── SETTINGS ────────────────────────────────────────────────────
EXCEL_NAME   = "RHTP SPA Commitments.xlsx"
OUTPUT_NAME  = "spa-table.json"
SEARCH_DEPTH = 2

COUNT_TITLE = "State Policy Action Commitments"
NONE_TEXT   = "None"

EXCLUDED_STATES = set()   # e.g. {"District of Columbia"}

TSF_STRIP_RE = re.compile(
    r"\bTechnical Score Factor\b[\s:\u2013-]*", re.I)
SPA_RE   = re.compile(r"\bSPA\b")
SPA_FULL = "State Policy Action"
CODE_RE  = re.compile(r"([A-Z])\.\s*(\d+)", re.I)

STATE_NAMES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona",
    "AR": "Arkansas", "CA": "California",
    "CO": "Colorado", "CT": "Connecticut",
    "DE": "Delaware", "DC": "District of Columbia",
    "FL": "Florida", "GA": "Georgia", "HI": "Hawaii",
    "ID": "Idaho", "IL": "Illinois", "IN": "Indiana",
    "IA": "Iowa", "KS": "Kansas", "KY": "Kentucky",
    "LA": "Louisiana", "ME": "Maine", "MD": "Maryland",
    "MA": "Massachusetts", "MI": "Michigan",
    "MN": "Minnesota", "MS": "Mississippi",
    "MO": "Missouri", "MT": "Montana",
    "NE": "Nebraska", "NV": "Nevada",
    "NH": "New Hampshire", "NJ": "New Jersey",
    "NM": "New Mexico", "NY": "New York",
    "NC": "North Carolina", "ND": "North Dakota",
    "OH": "Ohio", "OK": "Oklahoma", "OR": "Oregon",
    "PA": "Pennsylvania", "RI": "Rhode Island",
    "SC": "South Carolina", "SD": "South Dakota",
    "TN": "Tennessee", "TX": "Texas", "UT": "Utah",
    "VT": "Vermont", "VA": "Virginia",
    "WA": "Washington", "WV": "West Virginia",
    "WI": "Wisconsin", "WY": "Wyoming",
    "PR": "Puerto Rico", "GU": "Guam",
    "VI": "Virgin Islands", "AS": "American Samoa",
    "MP": "Northern Mariana Islands",
}
OFFICIAL_BY_LOWER = {n.lower(): n
                     for n in STATE_NAMES.values()}


# ── HELPERS ─────────────────────────────────────────────────────
def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()

def clean_path(text):
    """Tidy a typed or drag-and-dropped path."""
    return text.strip().strip("'\"").replace("\\ ", " ")

def norm_code(v):
    m = CODE_RE.search(clean(v))
    if not m:
        return ""
    return f"{m.group(1).upper()}.{m.group(2)}"

def code_sort_key(code):
    letter, _, num = code.partition(".")
    return (letter, int(num) if num.isdigit() else 0)

def tidy(text):
    """Drop 'Technical Score Factor', spell out SPA."""
    text = TSF_STRIP_RE.sub("", clean(text))
    text = SPA_RE.sub(SPA_FULL, text)
    return re.sub(r" {2,}", " ", text).strip()

def column_title(text):
    """DataTable title: no commas, single spaces."""
    text = tidy(text).replace(",", " ")
    return re.sub(r"\s{2,}", " ", text).strip()

def fix_state(raw):
    """Official spelling, or None if not a state."""
    text = clean(raw)
    if text.upper() in STATE_NAMES:
        return STATE_NAMES[text.upper()]
    if text.lower() in OFFICIAL_BY_LOWER:
        return OFFICIAL_BY_LOWER[text.lower()]
    match = difflib.get_close_matches(
        text.lower(), OFFICIAL_BY_LOWER.keys(),
        n=1, cutoff=0.8)
    return OFFICIAL_BY_LOWER[match[0]] if match else None

def yes_no(value):
    """Normalize to 'Yes'/'No'. Returns (text, ok)."""
    text = clean(value)
    if text == "":
        return "No", True
    low = text.lower()
    if low in ("yes", "y", "x", "true",
               "\u2713", "\u2714"):
        return "Yes", True
    if low in ("no", "n", "false", "n/a"):
        return "No", True
    return text, False


# ── FIND THE TABLE ──────────────────────────────────────────────
def find_columns(cells):
    """cells = lowercased header row. Returns dict or None."""
    def first(test):
        return next((i for i, c in enumerate(cells)
                     if c and test(c)), None)
    col = {
        "state":  first(lambda c: c == "state"),
        "factor": first(lambda c: "score factor" in c
                        or c == "factor"),
        "item":   first(lambda c: "item" in c
                        or "policy" in c),
        "commit": first(lambda c: "commit" in c
                        or "status" in c),
    }
    return None if None in col.values() else col

def locate_table(wb):
    for ws in wb.worksheets:
        rows = ws.iter_rows(min_row=1, max_row=20,
                            values_only=True)
        for i, row in enumerate(rows):
            cells = [clean(c).lower() for c in row]
            col = find_columns(cells)
            if col:
                return ws, i, col
    return None, None, None


# ── FILE DISCOVERY ──────────────────────────────────────────────
def search_dirs():
    dirs = []
    here = os.path.dirname(os.path.abspath(__file__))
    for d in (here, os.getcwd()):
        d = os.path.abspath(d)
        if d not in dirs:
            dirs.append(d)
    return dirs

def walk_files(extensions):
    seen = set()
    for root_dir in search_dirs():
        base = root_dir.rstrip(os.sep).count(os.sep)
        for root, subdirs, files in os.walk(root_dir):
            if root.count(os.sep) - base >= SEARCH_DEPTH:
                subdirs[:] = []
            subdirs[:] = [d for d in subdirs
                          if not d.startswith(".")]
            for name in files:
                path = os.path.abspath(
                    os.path.join(root, name))
                if path in seen or name.startswith("~$"):
                    continue
                if name.lower().endswith(extensions):
                    seen.add(path)
                    yield path

def choose(candidates, label):
    if not candidates:
        return None
    candidates.sort(key=os.path.getmtime, reverse=True)
    if len(candidates) == 1:
        return candidates[0]
    print(f"\nFound {len(candidates)} possible "
          f"{label} files:")
    for i, p in enumerate(candidates, 1):
        print(f"  {i}. {p}")
    if not sys.stdin.isatty():
        print(f"Using the most recent: {candidates[0]}")
        return candidates[0]
    while True:
        ans = input(f"Which {label}? "
                    f"[1-{len(candidates)}, Enter = 1]: ")
        ans = ans.strip()
        if not ans:
            return candidates[0]
        if ans.isdigit() and 1 <= int(ans) <= len(candidates):
            return candidates[int(ans) - 1]
        print("  Please enter a number from the list.")

def has_table(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws, _, _ = locate_table(wb)
        wb.close()
        return ws is not None
    except Exception:
        return False

def find_excel():
    files = list(walk_files((".xlsx", ".xlsm")))
    named = [p for p in files
             if os.path.basename(p).lower()
             == EXCEL_NAME.lower()]
    if named:
        return choose(named, f"'{EXCEL_NAME}'")
    return choose([p for p in files if has_table(p)],
                  "Excel")


# ── MAIN ────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(
        description="Build a DataTable JSON from the "
                    "State Policy Action spreadsheet.")
    ap.add_argument("-i", "--input",
                    help="Spreadsheet (optional)")
    ap.add_argument("-o", "--output",
                    help="Output JSON (optional)")
    args = ap.parse_args()

    excel_file = clean_path(args.input) if args.input \
                 else find_excel()
    if not excel_file or not os.path.exists(excel_file):
        print(f"ERROR: Could not find '{EXCEL_NAME}' in "
              f"{', '.join(search_dirs())} or subfolders.")
        sys.exit(1)

    if args.output:
        output_file = clean_path(args.output)
        if not output_file.lower().endswith(".json"):
            output_file += ".json"
    else:
        output_file = os.path.join(
            os.path.dirname(os.path.abspath(excel_file)),
            OUTPUT_NAME)

    print(f"\nExcel:  {excel_file}")
    print(f"Output: {output_file}\n")

    try:
        wb = openpyxl.load_workbook(excel_file,
                                    data_only=True)
    except Exception as e:
        print(f"ERROR: Could not open Excel file: {e}")
        sys.exit(1)

    ws, header_idx, col = locate_table(wb)
    if ws is None:
        print("ERROR: Couldn't find a header row with "
              "State / Score Factor / Policy Item / "
              "Commitment columns. First rows:")
        for sheet in wb.worksheets:
            print(f"\n=== {sheet.title} ===")
            for row in sheet.iter_rows(min_row=1, max_row=6,
                                       values_only=True):
                print(row)
        sys.exit(1)

    # ── Read rows ───────────────────────────────────────────
    # items_by_code: code -> [raw policy items, in order]
    items_by_code = OrderedDict()
    raw_values = {}     # (state, code, item) -> value
    states = set()
    fixed, skipped, odd_values = set(), set(), []

    rows = ws.iter_rows(min_row=header_idx + 2,
                        values_only=True)
    for row in rows:
        cells = list(row)

        def get(key):
            i = col[key]
            return cells[i] if i < len(cells) else None

        raw_state = clean(get("state"))
        code = norm_code(get("factor"))
        item = clean(get("item"))
        if not raw_state or not code or not item:
            continue

        state = fix_state(raw_state)
        if state is None:
            skipped.add(raw_state)
            continue
        if state != raw_state:
            fixed.add(f"{raw_state} -> {state}")
        if state in EXCLUDED_STATES:
            continue

        code_items = items_by_code.setdefault(code, [])
        if item not in code_items:
            code_items.append(item)

        text, ok = yes_no(get("commit"))
        if not ok:
            odd_values.append(
                f"{state} / {code} / {item}: {text!r}")
        raw_values[(state, code, item)] = text
        states.add(state)

    # ── Column titles (factor order, then sheet order) ──────
    columns = []    # [(code, raw_item, title)]
    used = {"State", COUNT_TITLE}
    for code in sorted(items_by_code, key=code_sort_key):
        for item in items_by_code[code]:
            title = column_title(item)
            if title in used:            # keep unique
                title = f"{code} {title}"
            used.add(title)
            columns.append((code, item, title))

    # ── Build records ───────────────────────────────────────
    records, missing, none_states = [], set(), []
    for state in sorted(states):
        values = OrderedDict()
        for code, item, title in columns:
            key = (state, code, item)
            if key not in raw_values:
                missing.add(state)
            values[title] = raw_values.get(key, "No")

        yes_count = sum(1 for v in values.values()
                        if v == "Yes")
        if yes_count == 0:
            none_states.append(state)

        record = OrderedDict()
        record["State"] = state
        record[COUNT_TITLE] = yes_count if yes_count \
                              else NONE_TEXT
        record.update(values)
        records.append(record)

    with open(output_file, "w", encoding="utf-8") as f:
        json.dump({"aaData": records}, f, indent=2,
                  ensure_ascii=False)

    # ── Report ──────────────────────────────────────────────
    print(f"Read tab '{ws.title}' "
          f"(header on row {header_idx + 1})")
    print(f"Wrote {len(records)} records to {output_file}")
    print(f"Columns: State, {COUNT_TITLE}, then "
          f"{len(columns)} policy columns:")
    for n, (code, _, title) in enumerate(columns, start=3):
        print(f"  {n:>2}. [{code}] {title}")
    if none_states:
        print(f"\nNo commitments ({COUNT_TITLE} = "
              f"{NONE_TEXT}): {', '.join(none_states)}")
    if fixed:
        print("Fixed state spellings: "
              + "; ".join(sorted(fixed)))
    if skipped:
        print("Skipped rows (not a state name): "
              + ", ".join(sorted(skipped)))
    if missing:
        print("Missing some policy rows (filled with "
              "'No'): " + ", ".join(sorted(missing)))
    for v in odd_values:
        print(f"WARNING - not Yes/No, kept as-is "
              f"(not counted): {v}")


if __name__ == "__main__":
    main()
